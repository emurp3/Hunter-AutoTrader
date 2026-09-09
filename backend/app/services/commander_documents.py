"""
Storage and retrieval for Commander-supplied reference documents —
capability profiles, resumes, and similar material Hunter can read from
when scoring opportunities or answering Commander in chat.

This is storage only. Per the standing addendum Commander-checkpoint
rule (see execution_accounting.py), having a document on file lets
Hunter reference and propose from it; any consequential external
submission built from it still requires its own Commander-approved
step — nothing here executes anything.
"""

from __future__ import annotations

import io
import logging
import os
import uuid
from pathlib import Path
from typing import Optional

from sqlmodel import Session, select

from app.models.commander_document import CommanderDocument

logger = logging.getLogger(__name__)

_SEED_PROFILE_PATH = Path(__file__).resolve().parent.parent / "data" / "commander_capability_profile.txt"


class UnsupportedDocumentType(ValueError):
    pass


def _documents_dir() -> Path:
    root = Path(os.getenv("HUNTER_DOCUMENTS_DIR", "/data/commander_documents"))
    root.mkdir(parents=True, exist_ok=True)
    return root


def _extract_text(filename: str, content_type: str, raw_bytes: bytes) -> str:
    lower_name = (filename or "").lower()
    if lower_name.endswith(".docx") or "wordprocessingml" in (content_type or ""):
        try:
            import docx  # python-docx
        except ImportError as exc:
            raise UnsupportedDocumentType(
                "DOCX support requires the python-docx package, which is not installed."
            ) from exc
        document = docx.Document(io.BytesIO(raw_bytes))
        return "\n".join(p.text for p in document.paragraphs)

    if lower_name.endswith(".pdf") or content_type == "application/pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise UnsupportedDocumentType(
                "PDF support requires the pypdf package, which is not installed."
            ) from exc
        reader = PdfReader(io.BytesIO(raw_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if lower_name.endswith((".txt", ".md")) or (content_type or "").startswith("text/"):
        return raw_bytes.decode("utf-8", errors="replace")

    raise UnsupportedDocumentType(
        f"Unsupported document type for '{filename}' ({content_type}). "
        "Supported: .docx, .pdf, .txt, .md"
    )


def save_document(
    session: Session,
    *,
    filename: str,
    content_type: str,
    raw_bytes: bytes,
    document_type: str = "other",
    source: str = "commander_upload",
) -> CommanderDocument:
    extracted_text = _extract_text(filename, content_type, raw_bytes)

    document_id = uuid.uuid4().hex
    stored_path: Optional[str] = None
    try:
        target = _documents_dir() / f"{document_id}_{filename}"
        target.write_bytes(raw_bytes)
        stored_path = str(target)
    except OSError as exc:
        # Storing the raw file is a convenience, not a requirement — the
        # extracted text (what Hunter actually reads) is what matters and
        # is always persisted in the DB row below.
        logger.warning("Could not write document to disk, keeping extracted text only: %s", exc)

    doc = CommanderDocument(
        document_id=document_id,
        filename=filename,
        document_type=document_type,
        content_type=content_type,
        stored_path=stored_path,
        extracted_text=extracted_text,
        source=source,
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc


def list_documents(session: Session) -> list[CommanderDocument]:
    return list(session.exec(select(CommanderDocument).order_by(CommanderDocument.uploaded_at.desc())).all())


def get_capability_profile_text(session: Session) -> Optional[str]:
    doc = session.exec(
        select(CommanderDocument)
        .where(CommanderDocument.document_type == "capability_profile")
        .order_by(CommanderDocument.uploaded_at.desc())
    ).first()
    return doc.extracted_text if doc else None


def seed_capability_profile(session: Session) -> Optional[CommanderDocument]:
    """Idempotent: seeds the bundled capability-profile text file as a
    CommanderDocument on first run only. Safe to call on every startup."""
    existing = session.exec(
        select(CommanderDocument).where(CommanderDocument.document_type == "capability_profile")
    ).first()
    if existing:
        return None

    if not _SEED_PROFILE_PATH.exists():
        logger.warning("Capability profile seed file not found at %s", _SEED_PROFILE_PATH)
        return None

    raw_text = _SEED_PROFILE_PATH.read_text(encoding="utf-8")
    doc = CommanderDocument(
        document_id=uuid.uuid4().hex,
        filename=_SEED_PROFILE_PATH.name,
        document_type="capability_profile",
        content_type="text/plain",
        stored_path=None,
        extracted_text=raw_text,
        source="seed",
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc
